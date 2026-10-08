import json
from datetime import datetime
from io import BytesIO
from flask import request, jsonify, current_app, session, send_file
from flask_login import login_required
from sqlalchemy.sql import text
from . import api_bp

def get_engine():
    return current_app.extensions['sqlalchemy'].engine

def date_to_key(date_str):
    if not date_str:
        return None
    return int(date_str.replace('-', ''))

# ------------------------------------------------------------------
# Report definitions (list shown in the main table)
# ------------------------------------------------------------------
REPORTS = [
    {"id": "production", "name_key": "reports.production", "desc_key": "reports.production_desc"},
    {"id": "inventory",  "name_key": "reports.inventory",  "desc_key": "reports.inventory_desc"},
    {"id": "quality",    "name_key": "reports.quality",    "desc_key": "reports.quality_desc"},
    {"id": "purchasing", "name_key": "reports.purchasing", "desc_key": "reports.purchasing_desc"},
]

@api_bp.route('/reports/list')
@login_required
def list_reports():
    """Return the list of available report definitions (for the main table)."""
    return jsonify(REPORTS)

# ------------------------------------------------------------------
# Run report – returns preview data (first 200 rows)
# ------------------------------------------------------------------

@api_bp.route('/reports/run', methods=['POST'])
@login_required
def run_report():
    data = request.get_json()
    report_type = data.get('report_type')
    date_from = data.get('date_from')
    date_to = data.get('date_to')
    
    if not report_type:
        return jsonify({'error': 'Missing report_type'}), 400
    
    from_key = date_to_key(date_from) if date_from else None
    to_key = date_to_key(date_to) if date_to else None
    
    engine = get_engine()
    with engine.connect() as conn:
        if report_type == 'production':
            sql = """
                SELECT 
                    po.ProductionOrderId,
                    po.ProductId,
                    p.descEnglish AS ProductName,
                    po.OrderQuantity,
                    po.Status,
                    po.StartDateKey,
                    po.EndDateKey,
                    po.ReleasedDateKey,
                    COALESCE(pod.ActualQuantity, 0) AS CompletedQuantity
                FROM t_ProductionOrder po
                JOIN t_Product p ON po.ProductId = p.ProductId
                LEFT JOIN t_ProductionOrderDetail pod ON po.ProductionOrderId = pod.ProductionOrderId
                WHERE 1=1
                """
            params = {}
            if from_key:
                sql += " AND po.StartDateKey >= :from_key"
                params['from_key'] = from_key
            if to_key:
                sql += " AND po.StartDateKey <= :to_key"
                params['to_key'] = to_key
            sql += " ORDER BY po.StartDateKey DESC LIMIT 200"
            rows = conn.execute(text(sql), params).fetchall()
            columns = ['ProductionOrderID', 'ProductID', 'ProductName', 'OrderQuantity', 'Status', 'StartDateKey', 'EndDateKey', 'ReleasedDateKey', 'CompletedQuantity']
            data_rows = [dict(zip(columns, row)) for row in rows]
            
        elif report_type == 'inventory':
            # Uses t_InventoryOnHand – sums all rows (approximate). For better accuracy, consider latest snapshot per product.
            sql = """
                SELECT 
                    p.ProductId,
                    p.descEnglish AS ProductName,
                    p.UOM,
                    COALESCE(SUM(ioh.Qty), 0) AS OnHand,
                    p.StandardCost,
                    p.SafetyStock
                FROM t_Product p
                LEFT JOIN t_InventoryOnHand ioh ON p.ProductId = ioh.ProductID
                GROUP BY p.ProductId
                ORDER BY p.ProductId LIMIT 200
                """
            rows = conn.execute(text(sql)).fetchall()
            columns = ['ProductID', 'ProductName', 'UOM', 'OnHand', 'StandardCost', 'SafetyStock']
            data_rows = [dict(zip(columns, row)) for row in rows]
            
        elif report_type == 'quality':
            sql = """
                SELECT 
                    nc.NCRID,
                    nc.NCRNumber,
                    nc.ProductID,
                    COALESCE(p.descEnglish, 'N/A') AS ProductName,
                    nc.Quantity,
                    nc.Severity,
                    nc.DefectCode,
                    nc.DefectDescription,
                    nc.Status,
                    CAST(strftime('%Y%m%d', nc.DateRaised) AS INTEGER) AS DateKey
                FROM t_NonConformance nc
                LEFT JOIN t_Product p ON nc.ProductID = p.ProductId
                WHERE 1=1
                """
            params = {}
            if from_key:
                sql += " AND CAST(strftime('%Y%m%d', nc.DateRaised) AS INTEGER) >= :from_key"
                params['from_key'] = from_key
            if to_key:
                sql += " AND CAST(strftime('%Y%m%d', nc.DateRaised) AS INTEGER) <= :to_key"
                params['to_key'] = to_key
            sql += " ORDER BY nc.DateRaised DESC LIMIT 200"
            rows = conn.execute(text(sql), params).fetchall()
            columns = ['NCRID', 'NCRNumber', 'ProductID', 'ProductName', 'Quantity', 'Severity', 'DefectCode', 'DefectDescription', 'Status', 'DateKey']
            data_rows = [dict(zip(columns, row)) for row in rows]
            
        elif report_type == 'purchasing':
            sql = """
                SELECT 
                    po.PurchaseOrderID,
                    s.SupplierName,
                    po.OrderDateKey,
                    po.ExpectedDeliveryDateKey,
                    (po.Qty * po.Unitprice) AS TotalAmount,
                    po.Status,
                    po.ProductID
                FROM t_PurchaseOrder po
                LEFT JOIN t_Suppliers s ON po.SupplierID = s.SupplierID
                WHERE 1=1
                """
            params = {}
            if from_key:
                sql += " AND po.OrderDateKey >= :from_key"
                params['from_key'] = from_key
            if to_key:
                sql += " AND po.OrderDateKey <= :to_key"
                params['to_key'] = to_key
            sql += " ORDER BY po.OrderDateKey DESC LIMIT 200"
            rows = conn.execute(text(sql), params).fetchall()
            columns = ['PurchaseOrderID', 'SupplierName', 'OrderDateKey', 'ExpectedDeliveryDateKey', 'TotalAmount', 'Status', 'ProductID']
            data_rows = [dict(zip(columns, row)) for row in rows]
        else:
            return jsonify({'error': 'Invalid report_type'}), 400
        
        session['last_report'] = {
            'report_type': report_type,
            'date_from': date_from,
            'date_to': date_to,
            'data': data_rows,
            'columns': columns
        }
        return jsonify({'columns': columns, 'data': data_rows})

# ------------------------------------------------------------------
# Export to Excel or PDF
# ------------------------------------------------------------------
@api_bp.route('/reports/export/<format>')
@login_required
def export_report(format):
    if format not in ['excel', 'pdf']:
        return jsonify({'error': 'Invalid format'}), 400
    
    last = session.get('last_report')
    if not last:
        return jsonify({'error': 'No report data. Please run the report first.'}), 400
    
    report_type = last['report_type']
    data = last['data']
    columns = last['columns']
    
    # Build a human-readable report title
    title_map = {
        'production': 'Production Report',
        'inventory': 'Inventory Report',
        'quality': 'Quality Report',
        'purchasing': 'Purchasing Report'
    }
    title = title_map.get(report_type, 'Operational Report')
    date_range = ""
    if last.get('date_from') or last.get('date_to'):
        date_range = f" ({last.get('date_from', '')} - {last.get('date_to', '')})"
    
    if format == 'excel':
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment
        except ImportError:
            return jsonify({'error': 'openpyxl not installed. Run: pip install openpyxl'}), 500
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = report_type.capitalize()
        
        # Title row
        ws.merge_cells('A1:{}1'.format(openpyxl.utils.get_column_letter(len(columns))))
        ws['A1'] = f"{title}{date_range}"
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')
        
        # Header row
        headers = [col.replace('Key', '').replace('ID', ' ID') for col in columns]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_idx, value=header)
            cell.font = Font(bold=True)
        
        # Data rows
        for row_idx, row in enumerate(data, 3):
            for col_idx, col_name in enumerate(columns, 1):
                value = row.get(col_name)
                # Convert integer date keys to readable dates
                if col_name.endswith('DateKey') and value and isinstance(value, int):
                    try:
                        value = datetime.strptime(str(value), '%Y%m%d').strftime('%Y-%m-%d')
                    except:
                        pass
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 2, 40)
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=f"{report_type}_report_{datetime.now().strftime('%Y%m%d')}.xlsx", mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    
    elif format == 'pdf':
        try:
            from reportlab.lib.pagesizes import landscape, A4
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib import colors
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
        except ImportError:
            return jsonify({'error': 'reportlab not installed. Run: pip install reportlab'}), 500
        
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        story = []
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], alignment=1, fontSize=16)
        story.append(Paragraph(f"{title}{date_range}", title_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Prepare table data
        table_data = [[col.replace('Key', '').replace('ID', ' ID') for col in columns]]
        for row in data:
            table_row = []
            for col in columns:
                val = row.get(col)
                if col.endswith('DateKey') and val and isinstance(val, int):
                    try:
                        val = datetime.strptime(str(val), '%Y%m%d').strftime('%Y-%m-%d')
                    except:
                        pass
                table_row.append(str(val) if val is not None else '')
            table_data.append(table_row)
        
        # Create table
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.grey),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('BACKGROUND', (0,1), (-1,-1), colors.beige),
            ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
            ('FONTSIZE', (0,1), (-1,-1), 8),
        ]))
        story.append(table)
        doc.build(story)
        buffer.seek(0)
        return send_file(buffer, as_attachment=True, download_name=f"{report_type}_report_{datetime.now().strftime('%Y%m%d')}.pdf", mimetype='application/pdf')